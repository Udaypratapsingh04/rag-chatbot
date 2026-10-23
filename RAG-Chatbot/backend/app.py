"""
RAG Chatbot backend (Flask).

Upload PDF / TXT / DOCX files, index them in a local FAISS vector store and
answer questions about them with Google Gemini or OpenAI.

API keys can come from either place (the Settings screen wins over .env):
  1. The Settings dialog in the web UI  -> saved to backend/settings.json
  2. backend/.env                       -> GOOGLE_API_KEY / OPENAI_API_KEY
"""

import json
import logging
import os
import re
import shutil
import threading
import time
import traceback
import uuid
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, request
from flask_cors import CORS
from werkzeug.exceptions import HTTPException
from werkzeug.utils import secure_filename

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

UPLOAD_DIR = BASE_DIR / "uploads"
VECTOR_DIR = BASE_DIR / "vector_store"
INDEX_DIR = VECTOR_DIR / "faiss_index"          # holds index.faiss + index.pkl
INFO_FILE = VECTOR_DIR / "index_info.json"      # which embedding model built the index
METADATA_FILE = VECTOR_DIR / "metadata.json"    # list of indexed documents
SETTINGS_FILE = BASE_DIR / "settings.json"      # keys / models saved from the UI

UPLOAD_DIR.mkdir(exist_ok=True)
VECTOR_DIR.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {"pdf", "txt", "docx"}
MAX_UPLOAD_MB = 25
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
TOP_K = 5

# Model lists are tried in order; the first one the API accepts is used. This
# means a model being retired no longer breaks the app. Users can also pin an
# exact model name in Settings.
PROVIDERS = {
    "gemini": {
        "label": "Google Gemini",
        "key_envs": ["GOOGLE_API_KEY", "GEMINI_API_KEY"],
        "model_envs": {"chat": "GEMINI_CHAT_MODEL", "embedding": "GEMINI_EMBEDDING_MODEL"},
        "chat_models": ["gemini-flash-latest", "gemini-2.5-flash", "gemini-3.5-flash"],
        "embedding_models": ["gemini-embedding-001", "gemini-embedding-2"],
        "batch_size": 32,
    },
    "openai": {
        "label": "OpenAI",
        "key_envs": ["OPENAI_API_KEY"],
        "model_envs": {"chat": "OPENAI_CHAT_MODEL", "embedding": "OPENAI_EMBEDDING_MODEL"},
        "chat_models": ["gpt-5-mini", "gpt-4.1-mini", "gpt-4o-mini"],
        "embedding_models": ["text-embedding-3-small", "text-embedding-3-large"],
        "batch_size": 64,
    },
}

SYSTEM_PROMPT = (
    "You answer questions about the user's uploaded documents. "
    "Use the provided context excerpts as your source of truth. "
    "If the answer is not in the context, say so plainly instead of guessing. "
    "Be concise and accurate, and mention which document an answer came from when it helps."
)

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
log = logging.getLogger("rag")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024
CORS(app, resources={r"/api/*": {"origins": [
    "http://localhost:5173", "http://127.0.0.1:5173",
]}})

STATE_LOCK = threading.RLock()   # guards the FAISS files, metadata and settings
_resolved_embedding = {}         # (provider, pinned) -> working embedding model
_resolved_chat = {}              # (provider, pinned) -> working chat model


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------

class AppError(Exception):
    """An error with a message that is safe and useful to show to the user."""

    def __init__(self, status, message, code=None, detail=None):
        super().__init__(message)
        self.status = status
        self.message = message
        self.code = code
        self.detail = detail


def classify(exc):
    """Roughly categorise an exception raised by a provider SDK."""
    low = str(exc).lower()

    def has(*subs):
        return any(s in low for s in subs)

    if has("insufficient_quota"):
        return "billing"
    if re.search(r"\b429\b", low) or has("resource_exhausted", "rate limit", "rate_limit", "quota"):
        return "quota"
    if has("api key not valid", "api_key_invalid", "invalid api key", "incorrect api key",
           "invalid_api_key", "unauthenticated", "permission_denied", "permission denied",
           "api key expired", "api keys are not supported") or re.search(r"\b(401|403)\b", low):
        return "auth"
    if re.search(r"\b404\b", low) or has("not found", "not_found", "model_not_found", "does not exist",
                                         "is not supported for", "no longer available",
                                         "has been shut down", "decommissioned"):
        return "model"
    if has("connection error", "connecterror", "connection refused", "timed out", "timeout",
           "name or service not known", "getaddrinfo", "temporary failure in name resolution",
           "max retries exceeded") or re.search(r"\bssl\b", low):
        return "network"
    return "other"


def to_app_error(exc, provider):
    """Convert a provider exception into an AppError with a friendly message."""
    if isinstance(exc, AppError):
        return exc
    label = PROVIDERS[provider]["label"]
    kind = classify(exc)
    detail = str(exc)[:500]
    log.warning("%s call failed (%s): %s", label, kind, detail)
    if kind == "auth":
        return AppError(401, f"{label} rejected the API key. Open Settings and check that the key is correct and active.", "auth", detail)
    if kind == "billing":
        return AppError(402, f"{label} says this account has no remaining credit. Check your billing settings.", "billing", detail)
    if kind == "quota":
        return AppError(429, f"{label} rate limit or quota reached. Wait a minute and try again, or check your plan limits.", "quota", detail)
    if kind == "model":
        return AppError(502, f"{label} could not find the requested model. Pick another model in Settings.", "model", detail)
    if kind == "network":
        return AppError(502, f"Could not reach {label}. Check your internet connection and try again.", "network", detail)
    return AppError(500, f"{label} returned an error: {detail[:200]}", "other", detail)


def parse_retry_delay(text):
    """Pull a 'retry in N seconds' hint out of a rate-limit error message."""
    m = re.search(r"retrydelay['\"]?\s*:\s*['\"]([0-9.]+)s", text, re.I)
    if m:
        return float(m.group(1))
    m = re.search(r"(?:retry|try again)[^0-9]{0,20}([0-9]+(?:\.[0-9]+)?)\s*(ms|s)", text, re.I)
    if m:
        value = float(m.group(1))
        return value / 1000 if m.group(2).lower() == "ms" else value
    return None


def with_retry(fn, attempts=5):
    """Call fn(), waiting and retrying when the provider says we're rate-limited."""
    delay = 2.0
    for i in range(attempts):
        try:
            return fn()
        except Exception as exc:  # noqa: BLE001
            if classify(exc) != "quota" or i == attempts - 1:
                raise
            wait = min(parse_retry_delay(str(exc)) or delay, 30.0) + 0.5
            log.info("Rate limited, retrying in %.1fs (attempt %d/%d)", wait, i + 1, attempts)
            time.sleep(wait)
            delay = min(delay * 2, 30.0)


# ---------------------------------------------------------------------------
# JSON helpers and settings
# ---------------------------------------------------------------------------

def read_json(path, default):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def write_json(path, data):
    """Write atomically so a crash never leaves a half-written file."""
    tmp = Path(str(path) + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    os.replace(tmp, path)


def load_settings():
    return read_json(SETTINGS_FILE, {})


def get_api_key(provider):
    """Return (key, source). The Settings screen wins over environment variables."""
    saved = (load_settings().get("api_keys") or {}).get(provider, "").strip()
    if saved:
        return saved, "settings"
    for name in PROVIDERS[provider]["key_envs"]:
        value = (os.getenv(name) or "").strip().strip('"').strip("'")
        if value:
            return value, "env"
    return None, None


def mask_key(key):
    if not key:
        return None
    return key[:3] + "…" + key[-4:] if len(key) > 10 else "••••"


def pinned_model(provider, kind):
    """A model name the user chose explicitly (Settings first, then .env), else ''."""
    saved = ((load_settings().get("models") or {}).get(provider) or {}).get(kind, "").strip()
    if saved:
        return saved
    return (os.getenv(PROVIDERS[provider]["model_envs"][kind]) or "").strip()


def active_provider():
    chosen = load_settings().get("provider") or (os.getenv("LLM_PROVIDER") or "").strip().lower()
    if chosen in PROVIDERS:
        return chosen
    # No explicit choice: use whichever provider has a key.
    for name in PROVIDERS:
        if get_api_key(name)[0]:
            return name
    return "gemini"


def require_key(provider):
    key, _ = get_api_key(provider)
    if not key:
        label = PROVIDERS[provider]["label"]
        raise AppError(400, f"No {label} API key yet. Open Settings and paste your key.", "no_key")
    return key


# ---------------------------------------------------------------------------
# LangChain wiring (imported lazily so the server can start and report
# problems clearly even when one provider's package is missing)
# ---------------------------------------------------------------------------

def make_embeddings(provider, model, key):
    if provider == "gemini":
        from langchain_google_genai import GoogleGenerativeAIEmbeddings
        return GoogleGenerativeAIEmbeddings(model=model, google_api_key=key)
    try:
        from langchain_openai import OpenAIEmbeddings
    except ImportError as exc:
        raise AppError(500, "OpenAI support is not installed. Run:  pip install langchain-openai", "missing_package") from exc
    return OpenAIEmbeddings(model=model, api_key=key)


def make_chat(provider, model, key):
    if provider == "gemini":
        from langchain_google_genai import ChatGoogleGenerativeAI
        return ChatGoogleGenerativeAI(model=model, google_api_key=key)
    try:
        from langchain_openai import ChatOpenAI
    except ImportError as exc:
        raise AppError(500, "OpenAI support is not installed. Run:  pip install langchain-openai", "missing_package") from exc
    return ChatOpenAI(model=model, api_key=key)


def get_faiss():
    from langchain_community.vectorstores import FAISS
    return FAISS


def resolve_embedding_model(provider, key):
    """Find an embedding model that works, trying known names in order."""
    pinned = pinned_model(provider, "embedding")
    cache_key = (provider, pinned)
    if cache_key in _resolved_embedding:
        return _resolved_embedding[cache_key]

    candidates = [pinned] if pinned else PROVIDERS[provider]["embedding_models"]
    tried = []
    for model in candidates:
        try:
            emb = make_embeddings(provider, model, key)
            with_retry(lambda: emb.embed_query("ping"))
            _resolved_embedding[cache_key] = model
            log.info("Using %s embedding model: %s", provider, model)
            return model
        except Exception as exc:  # noqa: BLE001
            if isinstance(exc, AppError):
                raise
            if classify(exc) == "model":
                tried.append(model)
                continue
            raise to_app_error(exc, provider) from exc
    label = PROVIDERS[provider]["label"]
    raise AppError(502, f"None of the {label} embedding models I tried are available ({', '.join(tried)}). "
                        "Enter a current model name in Settings.", "model")


def to_text(content):
    """LangChain message content may be a string or a list of content blocks."""
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        parts = []
        for block in content:
            if isinstance(block, str):
                parts.append(block)
            elif isinstance(block, dict) and block.get("type") in (None, "text") and block.get("text"):
                parts.append(block["text"])
        return "".join(parts)
    return str(content)


def run_chat(provider, key, messages):
    """Call the chat model, falling back through known model names if one is gone."""
    pinned = pinned_model(provider, "chat")
    cache_key = (provider, pinned)
    candidates = [pinned] if pinned else list(PROVIDERS[provider]["chat_models"])
    if cache_key in _resolved_chat and _resolved_chat[cache_key] in candidates:
        candidates.remove(_resolved_chat[cache_key])
        candidates.insert(0, _resolved_chat[cache_key])

    tried = []
    for model in candidates:
        try:
            llm = make_chat(provider, model, key)
            response = with_retry(lambda: llm.invoke(messages))
            _resolved_chat[cache_key] = model
            return to_text(response.content).strip(), model
        except Exception as exc:  # noqa: BLE001
            if isinstance(exc, AppError):
                raise
            if classify(exc) == "model":
                tried.append(model)
                continue
            raise to_app_error(exc, provider) from exc
    label = PROVIDERS[provider]["label"]
    raise AppError(502, f"None of the {label} chat models I tried are available ({', '.join(tried)}). "
                        "Enter a current model name in Settings.", "model")


# ---------------------------------------------------------------------------
# Document loading
# ---------------------------------------------------------------------------

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def safe_name(original):
    """secure_filename() erases non-ASCII names completely, so keep the extension and invent a stem."""
    ext = original.rsplit(".", 1)[1].lower()
    stem = secure_filename(Path(original).stem) or f"document-{uuid.uuid4().hex[:6]}"
    return f"{stem}.{ext}"


def read_pdf(path):
    from pypdf import PdfReader
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise AppError(422, "This PDF is password-protected. Remove the password and upload it again.", "pdf_encrypted")
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001 - a single bad page shouldn't sink the file
                text = ""
            if text.strip():
                pages.append((number, text))
    except AppError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise AppError(422, "This PDF could not be read. It may be damaged.", "pdf_unreadable", str(exc)) from exc
    if not pages:
        raise AppError(422, "No selectable text found. This PDF looks like a scan made of images. "
                            "Run OCR on it first (for example with ocrmypdf), then upload it again.", "pdf_no_text")
    return pages


def read_docx(path):
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
    except Exception as exc:  # noqa: BLE001
        raise AppError(422, "This Word file could not be read. It may be damaged or an old .doc file.", "docx_unreadable", str(exc)) from exc
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return [(None, "\n".join(lines))]


def read_txt(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return [(None, raw.decode(encoding))]
        except UnicodeError:
            continue
    return [(None, raw.decode("utf-8", errors="replace"))]


def load_chunks(path, filename):
    """Read a file and split it into LangChain Documents ready for embedding."""
    from langchain_core.documents import Document
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    ext = filename.rsplit(".", 1)[1].lower()
    pieces = {"pdf": read_pdf, "docx": read_docx, "txt": read_txt}[ext](path)

    raw_docs = []
    for page, text in pieces:
        if not text.strip():
            continue
        meta = {"source": filename}
        if page is not None:
            meta["page"] = page
        raw_docs.append(Document(page_content=text, metadata=meta))
    if not raw_docs:
        raise AppError(422, "No text could be extracted from this file.", "empty")

    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    return splitter.split_documents(raw_docs)


# ---------------------------------------------------------------------------
# Vector store
# ---------------------------------------------------------------------------

def index_exists():
    return (INDEX_DIR / "index.faiss").exists()


def load_index(embeddings):
    if not index_exists():
        return None
    return get_faiss().load_local(str(INDEX_DIR), embeddings, allow_dangerous_deserialization=True)


def build_index(docs, embeddings, batch_size):
    """Embed docs in small batches (with rate-limit retries). Returns (db, ids)."""
    FAISS = get_faiss()
    db, all_ids = None, []
    for start in range(0, len(docs), batch_size):
        batch = docs[start:start + batch_size]
        ids = [uuid.uuid4().hex for _ in batch]
        if db is None:
            db = with_retry(lambda: FAISS.from_documents(batch, embeddings, ids=ids))
        else:
            with_retry(lambda: db.add_documents(batch, ids=ids))
        all_ids.extend(ids)
    return db, all_ids


def wipe_index():
    shutil.rmtree(INDEX_DIR, ignore_errors=True)
    INFO_FILE.unlink(missing_ok=True)


def index_matches(provider, model):
    info = read_json(INFO_FILE, None)
    return bool(info) and info.get("provider") == provider and info.get("model") == model


def stored_files():
    return sorted(p.name for p in UPLOAD_DIR.iterdir() if p.is_file() and allowed_file(p.name))


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.get("/api/health")
def health():
    provider = active_provider()
    return jsonify({"status": "ok", "provider": provider, "ready": bool(get_api_key(provider)[0])})


@app.get("/api/settings")
def get_settings():
    saved = load_settings()
    providers = {}
    for name, cfg in PROVIDERS.items():
        key, source = get_api_key(name)
        providers[name] = {
            "label": cfg["label"],
            "has_key": bool(key),
            "key_hint": mask_key(key),
            "key_source": source,
            "chat_model": ((saved.get("models") or {}).get(name) or {}).get("chat", ""),
            "embedding_model": ((saved.get("models") or {}).get(name) or {}).get("embedding", ""),
            "default_chat_models": cfg["chat_models"],
            "default_embedding_models": cfg["embedding_models"],
        }
    provider = active_provider()
    return jsonify({"provider": provider, "ready": providers[provider]["has_key"], "providers": providers})


@app.post("/api/settings")
def save_settings():
    body = request.get_json(silent=True) or {}
    with STATE_LOCK:
        saved = load_settings()
        saved.setdefault("api_keys", {})
        saved.setdefault("models", {})

        provider = body.get("provider")
        if provider is not None:
            if provider not in PROVIDERS:
                raise AppError(400, "Unknown provider.")
            saved["provider"] = provider

        for name, value in (body.get("api_keys") or {}).items():
            if name in PROVIDERS and isinstance(value, str) and value.strip():
                saved["api_keys"][name] = value.strip()
        for name in body.get("remove_keys") or []:
            saved["api_keys"].pop(name, None)

        for name, models in (body.get("models") or {}).items():
            if name in PROVIDERS and isinstance(models, dict):
                saved["models"][name] = {
                    "chat": str(models.get("chat", "")).strip(),
                    "embedding": str(models.get("embedding", "")).strip(),
                }

        write_json(SETTINGS_FILE, saved)
        _resolved_embedding.clear()
        _resolved_chat.clear()
    return get_settings()


@app.post("/api/settings/test")
def test_settings():
    """Check the key works and report which models will be used."""
    provider = active_provider()
    key = require_key(provider)
    try:
        embedding_model = resolve_embedding_model(provider, key)
        reply, chat_model = run_chat(provider, key, [("human", "Reply with the single word: ok")])
    except AppError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise to_app_error(exc, provider) from exc
    return jsonify({"ok": True, "provider": provider, "embedding_model": embedding_model,
                    "chat_model": chat_model, "reply": reply[:40]})


def documents_status():
    provider = active_provider()
    metadata = read_json(METADATA_FILE, [])
    known = {m["filename"] for m in metadata}
    unindexed = [f for f in stored_files() if f not in known]
    info = read_json(INFO_FILE, None)

    reason = None
    if info and metadata:
        pinned = pinned_model(provider, "embedding")
        if info.get("provider") != provider or (pinned and info.get("model") != pinned):
            reason = "model_changed"
    if reason is None and unindexed:
        reason = "unindexed_files"

    return {
        "documents": [{k: v for k, v in m.items() if k != "ids"} for m in metadata],
        "needs_reindex": reason is not None,
        "reindex_reason": reason,
        "unindexed_files": unindexed,
        "index_provider": info.get("provider") if info else None,
        "index_model": info.get("model") if info else None,
    }


@app.get("/api/documents")
def list_documents():
    with STATE_LOCK:
        return jsonify(documents_status())


@app.post("/api/upload")
def upload_document():
    file = request.files.get("file")
    if file is None:
        raise AppError(400, "No file was sent with the request.")
    if not file.filename:
        raise AppError(400, "No file selected.")
    if not allowed_file(file.filename):
        raise AppError(400, "Unsupported file type. Upload a PDF, TXT or DOCX file.")

    provider = active_provider()
    key = require_key(provider)
    filename = safe_name(file.filename)

    # Work on a temporary copy; only move it into uploads/ once indexing succeeds.
    tmp_path = UPLOAD_DIR / f".incoming-{uuid.uuid4().hex}-{filename}"
    try:
        file.save(tmp_path)
        if tmp_path.stat().st_size == 0:
            raise AppError(422, "The file is empty.", "empty")

        docs = load_chunks(tmp_path, filename)
        model = resolve_embedding_model(provider, key)

        with STATE_LOCK:
            if index_exists() and not index_matches(provider, model):
                info = read_json(INFO_FILE, {}) or {}
                raise AppError(409, f"Your documents were indexed with {info.get('model', 'another model')}, "
                                    f"but you are now using {model}. Re-index your documents to continue.", "reindex_required")

        embeddings = make_embeddings(provider, model, key)
        try:
            new_db, new_ids = build_index(docs, embeddings, PROVIDERS[provider]["batch_size"])
        except Exception as exc:  # noqa: BLE001
            raise to_app_error(exc, provider) from exc

        with STATE_LOCK:
            metadata = read_json(METADATA_FILE, [])
            existing = load_index(embeddings)

            # Re-uploading the same filename replaces its old chunks.
            previous = next((m for m in metadata if m["filename"] == filename), None)
            if existing is not None and previous and previous.get("ids"):
                try:
                    existing.delete(previous["ids"])
                except Exception as exc:  # noqa: BLE001
                    raise AppError(409, f"'{filename}' is already indexed and could not be replaced. "
                                        "Rename the file or clear all documents first.", "duplicate", str(exc)) from exc

            if existing is not None:
                existing.merge_from(new_db)
                final_db = existing
            else:
                final_db = new_db

            INDEX_DIR.mkdir(parents=True, exist_ok=True)
            final_db.save_local(str(INDEX_DIR))
            os.replace(tmp_path, UPLOAD_DIR / filename)

            metadata = [m for m in metadata if m["filename"] != filename]
            metadata.append({
                "filename": filename,
                "chunks": len(docs),
                "size": (UPLOAD_DIR / filename).stat().st_size,
                "ids": new_ids,
                "uploaded_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
            })
            write_json(METADATA_FILE, metadata)
            write_json(INFO_FILE, {"provider": provider, "model": model})

        return jsonify({"message": f"'{filename}' was indexed.", "filename": filename, "chunks": len(docs)})
    finally:
        tmp_path.unlink(missing_ok=True)


@app.post("/api/reindex")
def reindex():
    """Rebuild the whole index from the files in uploads/ with the current embedding model."""
    provider = active_provider()
    key = require_key(provider)
    files = stored_files()
    if not files:
        raise AppError(400, "There are no uploaded files to index.")

    model = resolve_embedding_model(provider, key)
    embeddings = make_embeddings(provider, model, key)

    all_docs, per_file, skipped = [], [], []
    for name in files:
        try:
            chunks = load_chunks(UPLOAD_DIR / name, name)
        except AppError as exc:
            skipped.append({"filename": name, "reason": exc.message})
            continue
        per_file.append((name, len(chunks)))
        all_docs.extend(chunks)
    if not all_docs:
        raise AppError(422, "None of the uploaded files contained readable text.", "empty")

    try:
        db, ids = build_index(all_docs, embeddings, PROVIDERS[provider]["batch_size"])
    except Exception as exc:  # noqa: BLE001
        raise to_app_error(exc, provider) from exc

    with STATE_LOCK:
        INDEX_DIR.mkdir(parents=True, exist_ok=True)
        db.save_local(str(INDEX_DIR))
        metadata, cursor = [], 0
        for name, count in per_file:
            metadata.append({
                "filename": name, "chunks": count,
                "size": (UPLOAD_DIR / name).stat().st_size,
                "ids": ids[cursor:cursor + count],
                "uploaded_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
            })
            cursor += count
        write_json(METADATA_FILE, metadata)
        write_json(INFO_FILE, {"provider": provider, "model": model})
    return jsonify({"message": f"Re-indexed {len(per_file)} document(s).", "indexed": len(per_file), "skipped": skipped})


@app.delete("/api/documents")
def clear_documents():
    with STATE_LOCK:
        for item in UPLOAD_DIR.iterdir():
            if item.is_file() and item.name != ".gitkeep":
                item.unlink(missing_ok=True)
        wipe_index()
        METADATA_FILE.unlink(missing_ok=True)
    return jsonify({"message": "All documents were removed."})


@app.delete("/api/documents/<path:filename>")
def delete_document(filename):
    with STATE_LOCK:
        metadata = read_json(METADATA_FILE, [])
        record = next((m for m in metadata if m["filename"] == filename), None)
        if record is None:
            raise AppError(404, "That document is not in the index.")

        remaining = [m for m in metadata if m["filename"] != filename]
        if not remaining:
            wipe_index()
        else:
            info = read_json(INFO_FILE, {}) or {}
            provider = info.get("provider") or active_provider()
            key = require_key(provider)
            embeddings = make_embeddings(provider, info["model"], key)
            db = load_index(embeddings)
            if db is not None and record.get("ids"):
                try:
                    db.delete(record["ids"])
                    db.save_local(str(INDEX_DIR))
                except Exception as exc:  # noqa: BLE001
                    raise AppError(500, "Could not remove this document from the index. Use Clear all and re-upload.",
                                   "delete_failed", str(exc)) from exc

        (UPLOAD_DIR / secure_filename(filename)).unlink(missing_ok=True)
        (UPLOAD_DIR / filename).unlink(missing_ok=True)
        if remaining:
            write_json(METADATA_FILE, remaining)
        else:
            METADATA_FILE.unlink(missing_ok=True)
    return jsonify({"message": f"'{filename}' was removed."})


@app.post("/api/chat")
def chat():
    data = request.get_json(silent=True) or {}
    user_message = str(data.get("message", "")).strip()
    history = data.get("conversation_history") or []
    if not user_message:
        raise AppError(400, "Message cannot be empty.")

    provider = active_provider()
    key = require_key(provider)

    with STATE_LOCK:
        has_docs = index_exists() and bool(read_json(METADATA_FILE, []))
        info = read_json(INFO_FILE, None)
    if not has_docs:
        return jsonify({"response": "There are no documents to search yet. Upload a PDF, TXT or DOCX file "
                                    "from the sidebar and I can answer questions about it.",
                        "sources": []})

    # Query the index with the model that built it, even if settings changed since.
    index_provider = (info or {}).get("provider", provider)
    index_model = (info or {}).get("model")
    index_key, _ = get_api_key(index_provider)
    if index_provider != provider or not index_key or not index_model:
        raise AppError(409, "Your documents were indexed with a different provider. Re-index them to continue.",
                       "reindex_required")

    # Short follow-ups ("and the second one?") retrieve better with the previous question attached.
    past_questions = [str(t.get("content", "")) for t in history if t.get("role") == "user"]
    search_query = " ".join(past_questions[-1:] + [user_message])

    try:
        embeddings = make_embeddings(index_provider, index_model, index_key)
        with STATE_LOCK:
            db = load_index(embeddings)
        results = with_retry(lambda: db.similarity_search(search_query, k=TOP_K)) if db else []
    except Exception as exc:  # noqa: BLE001
        raise to_app_error(exc, provider) from exc

    sources, context_parts = {}, []
    for doc in results:
        name = doc.metadata.get("source", "Unknown")
        page = doc.metadata.get("page")
        pages = sources.setdefault(name, [])
        if page is not None and page not in pages:
            pages.append(page)
        context_parts.append(f"[{name}{f', page {page}' if page is not None else ''}]\n{doc.page_content}")
    context = "\n\n---\n\n".join(context_parts)

    system = SYSTEM_PROMPT + (f"\n\nContext excerpts:\n\n{context}" if context else "")
    messages = [("system", system)]
    for turn in history[-10:]:
        role, content = turn.get("role"), str(turn.get("content", ""))
        if content and role in ("user", "assistant"):
            messages.append(("human" if role == "user" else "ai", content))
    messages.append(("human", user_message))

    answer, model_used = run_chat(provider, key, messages)
    return jsonify({
        "response": answer or "The model returned an empty answer. Try rephrasing your question.",
        "sources": [{"filename": n, "pages": sorted(p)} for n, p in sources.items()],
        "model": model_used,
    })


# ---------------------------------------------------------------------------
# Error handling: always answer with JSON so the UI can show a real message
# ---------------------------------------------------------------------------

@app.errorhandler(AppError)
def handle_app_error(err):
    body = {"error": err.message}
    if err.code:
        body["code"] = err.code
    if err.detail:
        body["detail"] = err.detail
    return jsonify(body), err.status


@app.errorhandler(HTTPException)
def handle_http_error(err):
    if err.code == 413:
        return jsonify({"error": f"That file is too large. The limit is {MAX_UPLOAD_MB} MB.", "code": "too_large"}), 413
    return jsonify({"error": err.description or err.name}), err.code or 500


@app.errorhandler(Exception)
def handle_unexpected(err):
    log.error("Unhandled error:\n%s", traceback.format_exc())
    return jsonify({"error": "Something went wrong on the server.", "detail": f"{type(err).__name__}: {err}"[:500],
                    "code": "server_error"}), 500


if __name__ == "__main__":
    # The reloader is off on purpose: a restart in the middle of an upload would drop the request.
    app.run(debug=False, host="127.0.0.1", port=5000, threaded=True)
